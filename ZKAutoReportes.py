import pandas
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path
import re
import sys

# Carpetas
INPUT_DIR  = Path("ZKEntrada")
OUTPUT_DIR = Path("ZKSalida")

# Formato del csv
SEPARADOR  = ";"
ENCODING   = "latin-1"
FILA_DATOS = 1
COL_NOMBRE = "Nombre"
COL_ID     = "ID"


COLOR_HEADER_BG = "2F5496"
COLOR_HEADER_FG = "FFFFFF"
COLOR_FILA_PAR  = "DCE6F1"

def limpiar_nombre_hoja(nombre: str) -> str:
    return re.sub(r"[\\/*?:\[\]]", "_", nombre)[:31]

def aplicar_estilo_encabezado(ws, n_cols: int):
    header_font  = Font(name="Arial", bold=True, color=COLOR_HEADER_FG, size=10)
    header_fill  = PatternFill("solid", fgColor=COLOR_HEADER_BG)
    header_align = Alignment(horizontal="center", vertical="center")
    thin   = Side(style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = header_align
        cell.border    = border
    ws.row_dimensions[1].height = 18

def aplicar_estilo_datos(ws, n_filas: int, n_cols: int):
    fill_par  = PatternFill("solid", fgColor=COLOR_FILA_PAR)
    data_font = Font(name="Arial", size=9)
    thin   = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in range(2, n_filas + 2):
        for col in range(1, n_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font   = data_font
            cell.border = border
            if row % 2 == 0:
                cell.fill = fill_par

def ajustar_columnas(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

def procesar_csv(csv_path: Path):
    sucursal = csv_path.stem
    print(f"\n-> Procesando: {csv_path.name}")

    df = pandas.read_csv(
        csv_path,
        sep=SEPARADOR,
        skiprows=FILA_DATOS,
        encoding=ENCODING,
        dtype=str
    )

    # Limpiar encabezado
    df.columns = df.columns.str.strip()
    df = df.dropna(how="all")
    #quitar col vacías
    df = df.loc[:, ~df.columns.str.startswith("Unnamed")]
    df[COL_NOMBRE] = df[COL_NOMBRE].str.strip()
    df[COL_ID]     = df[COL_ID].str.strip()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for nombre in sorted(df[COL_NOMBRE].dropna().unique()):
        df_persona = df[df[COL_NOMBRE] == nombre].reset_index(drop=True)
        ws = wb.create_sheet(title=limpiar_nombre_hoja(nombre))

        ws.append(df.columns.tolist())
        for _, fila in df_persona.iterrows():
            ws.append(fila.tolist())

        n_filas = len(df_persona)
        n_cols  = len(df.columns)

        aplicar_estilo_encabezado(ws, n_cols)
        aplicar_estilo_datos(ws, n_filas, n_cols)
        ajustar_columnas(ws)
        ws.freeze_panes = "A2"

        print(f"  OK {nombre} ({n_filas} registros)")

    salida = OUTPUT_DIR / f"{sucursal}.xlsx"
    wb.save(salida)
    print(f"  Guardado: {salida}")

def main():
    INPUT_DIR.mkdir(exist_ok=True)
    OUTPUT_DIR.mkdir(exist_ok=True)

    csvs = list(INPUT_DIR.glob("*.csv"))

    if not csvs:
        print(".csv file not found in ZKentrada folder")
        input("\nPresiona Enter para cerrar...")
        sys.exit(0)

    print(f"Encontrados {len(csvs)} archivo(s) CSV")

    for csv_path in csvs:
        try:
            procesar_csv(csv_path)
        except Exception as e:
            print(f" Error en {csv_path.name}: {e}")

    print(f"\nListo! Los Excel estan en la carpeta 'ZKsalida/'")
    input("\nPresiona Enter para cerrar...")

if __name__ == "__main__":
    main()